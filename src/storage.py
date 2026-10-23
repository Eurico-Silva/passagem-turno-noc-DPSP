"""Persistência Excel com bloqueio de arquivo, backup diário e histórico de eventos.

Excel é usado como armazenamento inicial para um piloto controlado. Para múltiplos
processos/instâncias ou alta concorrência, migrar para banco de dados transacional.
"""
from __future__ import annotations

import json
import os
import shutil
import tempfile
import uuid
from datetime import date, datetime, time
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

from filelock import FileLock, Timeout
from openpyxl import Workbook, load_workbook

from .config import DEFAULT_DATA_PATH, SHEETS, STATUS_TERMINAIS, TEMPLATE_PATH

TZ = ZoneInfo("America/Sao_Paulo")


def now_local() -> datetime:
    """Retorna horário local de São Paulo sem tzinfo, compatível com células Excel."""
    return datetime.now(TZ).replace(tzinfo=None, microsecond=0)


def combine_local(day: date | None, clock: time | None) -> datetime | None:
    if not day or not clock:
        return None
    return datetime.combine(day, clock).replace(second=0, microsecond=0)


def _plain(value: Any) -> Any:
    if isinstance(value, (datetime, date, time)):
        return value.isoformat(sep=" ") if isinstance(value, datetime) else value.isoformat()
    if value is None:
        return ""
    return value


def _safe_excel_value(value: Any) -> Any:
    """Evita que texto informado pelo usuário seja interpretado como fórmula Excel."""
    if isinstance(value, str) and value[:1] in {"=", "+", "-", "@"}:
        return "'" + value
    return value


def _same_value(a: Any, b: Any) -> bool:
    if a is None and (b is None or b == ""):
        return True
    if b is None and (a is None or a == ""):
        return True
    if isinstance(a, datetime) and isinstance(b, datetime):
        return a.replace(microsecond=0) == b.replace(microsecond=0)
    return str(a or "") == str(b or "")


class ExcelStore:
    def __init__(self, path: str | Path = DEFAULT_DATA_PATH, template_path: str | Path | None = TEMPLATE_PATH):
        self.path = Path(path)
        self.template_path = Path(template_path) if template_path else None
        self.lock_path = Path(str(self.path) + ".lock")
        self.backup_dir = self.path.parent / "backups"
        self.path.parent.mkdir(parents=True, exist_ok=True)
        self._initialize_if_missing()

    def _initialize_if_missing(self) -> None:
        if self.path.exists():
            return
        with FileLock(str(self.lock_path), timeout=10):
            if self.path.exists():
                return
            if self.template_path and self.template_path.exists():
                shutil.copy2(self.template_path, self.path)
                # Ensure headers/schema are aligned even if the template is older.
                wb = load_workbook(self.path)
                for sheet_name, columns in SHEETS.items():
                    if sheet_name not in wb.sheetnames:
                        ws = wb.create_sheet(sheet_name)
                        ws.append(columns)
                    else:
                        ws = wb[sheet_name]
                        if ws.max_row == 1 and ws.max_column == 1 and ws.cell(1, 1).value is None:
                            ws.append(columns)
                wb.save(self.path)
                wb.close()
            else:
                wb = Workbook()
                default = wb.active
                wb.remove(default)
                for sheet_name, columns in SHEETS.items():
                    ws = wb.create_sheet(sheet_name)
                    ws.append(columns)
                    ws.freeze_panes = "A2"
                    ws.auto_filter.ref = f"A1:{ws.cell(1, len(columns)).column_letter}1"
                wb.save(self.path)
                wb.close()

    def _load(self):
        if not self.path.exists():
            self._initialize_if_missing()
        return load_workbook(self.path)

    def _backup_once_per_day(self) -> None:
        if not self.path.exists():
            return
        self.backup_dir.mkdir(parents=True, exist_ok=True)
        backup = self.backup_dir / f"{self.path.stem}_{now_local():%Y%m%d}.xlsx"
        if not backup.exists():
            shutil.copy2(self.path, backup)

    def _atomic_save(self, wb) -> None:
        temp_name = None
        try:
            with tempfile.NamedTemporaryFile(prefix="passagem_", suffix=".xlsx", dir=self.path.parent, delete=False) as tmp:
                temp_name = tmp.name
            wb.save(temp_name)
            os.replace(temp_name, self.path)
        finally:
            if temp_name and os.path.exists(temp_name):
                os.unlink(temp_name)

    def _row_map(self, ws, row_number: int) -> dict[str, Any]:
        headers = [cell.value for cell in ws[1]]
        return {header: ws.cell(row_number, col).value for col, header in enumerate(headers, 1) if header}

    def _find_ticket_row(self, ws, ticket_id: str) -> int | None:
        needle = str(ticket_id).strip().casefold()
        for row_number in range(2, ws.max_row + 1):
            value = ws.cell(row_number, 1).value
            if value is not None and str(value).strip().casefold() == needle:
                return row_number
        return None

    def _append_history(self, wb, ticket_id: str, actor: str, event_type: str,
                        field: str = "", old_value: Any = "", new_value: Any = "",
                        comment: str = "") -> None:
        ws = wb["Historico"]
        ws.append([
            f"EVT-{uuid.uuid4().hex[:12].upper()}", _safe_excel_value(str(ticket_id)), now_local(),
            _safe_excel_value(actor), _safe_excel_value(event_type), _safe_excel_value(field),
            _safe_excel_value(str(_plain(old_value)) if old_value is not None else ""),
            _safe_excel_value(str(_plain(new_value)) if new_value is not None else ""),
            _safe_excel_value(comment),
        ])

    def list_records(self, sheet_name: str) -> list[dict[str, Any]]:
        if sheet_name not in SHEETS:
            raise ValueError(f"Aba não permitida: {sheet_name}")
        try:
            with FileLock(str(self.lock_path), timeout=10):
                wb = self._load()
                ws = wb[sheet_name]
                records = [self._row_map(ws, r) for r in range(2, ws.max_row + 1)]
                wb.close()
                return [r for r in records if any(v is not None and v != "" for v in r.values())]
        except Timeout as exc:
            raise RuntimeError("A base está ocupada por outra gravação. Tente novamente em alguns segundos.") from exc

    def create_ticket(self, data: dict[str, Any], operator: str) -> None:
        ticket_id = str(data.get("ticket_id", "")).strip()
        actor = str(operator).strip()
        if not ticket_id:
            raise ValueError("Informe o ID do chamado.")
        if not actor:
            raise ValueError("Informe o operador responsável pelo registro.")
        allowed = set(SHEETS["Chamados"])
        normalized = {k: v for k, v in data.items() if k in allowed}
        normalized["ticket_id"] = ticket_id
        now = now_local()
        normalized.setdefault("data_abertura", now)
        normalized.setdefault("status", "Aberto")
        normalized.setdefault("turno_atual", "Manhã")
        normalized["criado_em"] = now
        normalized["criado_por"] = actor
        normalized["atualizado_em"] = now
        normalized["atualizado_por"] = actor
        if normalized.get("status") in STATUS_TERMINAIS:
            normalized["encerrado_em"] = now
        else:
            normalized["encerrado_em"] = None
        try:
            with FileLock(str(self.lock_path), timeout=10):
                wb = self._load()
                ws = wb["Chamados"]
                if self._find_ticket_row(ws, ticket_id):
                    wb.close()
                    raise ValueError(f"O chamado {ticket_id} já existe. Use a tela de atualização.")
                self._backup_once_per_day()
                columns = SHEETS["Chamados"]
                ws.append([_safe_excel_value(normalized.get(c)) for c in columns])
                snapshot = {k: _plain(v) for k, v in normalized.items() if k in {
                    "loja", "tecnico", "prioridade", "status", "motivo_pendencia", "prazo_sla",
                    "ultima_atualizacao", "proxima_acao", "responsavel_proxima_acao", "turno_atual"
                }}
                self._append_history(wb, ticket_id, actor, "CRIACAO", "registro", "", json.dumps(snapshot, ensure_ascii=False), "Registro inicial criado.")
                self._atomic_save(wb)
                wb.close()
        except Timeout as exc:
            raise RuntimeError("A base está ocupada por outra gravação. Tente novamente em alguns segundos.") from exc

    def update_ticket(self, ticket_id: str, updates: dict[str, Any], operator: str, comment: str = "") -> int:
        actor = str(operator).strip()
        if not actor:
            raise ValueError("Informe o operador responsável pela atualização.")
        editable = {
            "loja", "tecnico", "prioridade", "status", "motivo_pendencia", "prazo_sla",
            "ultima_atualizacao", "proxima_acao", "responsavel_proxima_acao", "turno_atual", "resolucao"
        }
        updates = {k: v for k, v in updates.items() if k in editable}
        try:
            with FileLock(str(self.lock_path), timeout=10):
                wb = self._load()
                ws = wb["Chamados"]
                row = self._find_ticket_row(ws, ticket_id)
                if not row:
                    wb.close()
                    raise ValueError(f"Chamado {ticket_id} não encontrado.")
                current = self._row_map(ws, row)
                changes: list[tuple[str, Any, Any]] = []
                for field, new_value in updates.items():
                    old_value = current.get(field)
                    if not _same_value(old_value, new_value):
                        changes.append((field, old_value, new_value))
                if not changes:
                    wb.close()
                    return 0
                self._backup_once_per_day()
                now = now_local()
                for field, old_value, new_value in changes:
                    col = SHEETS["Chamados"].index(field) + 1
                    ws.cell(row, col).value = _safe_excel_value(new_value)
                    self._append_history(wb, ticket_id, actor, "ATUALIZACAO", field, old_value, new_value, comment)
                old_status = current.get("status")
                new_status = updates.get("status", old_status)
                if new_status in STATUS_TERMINAIS and old_status not in STATUS_TERMINAIS:
                    old_closed = current.get("encerrado_em")
                    new_closed = now
                    ws.cell(row, SHEETS["Chamados"].index("encerrado_em") + 1).value = new_closed
                    self._append_history(wb, ticket_id, actor, "ATUALIZACAO", "encerrado_em", old_closed, new_closed, "Encerramento automático por mudança de status.")
                elif old_status in STATUS_TERMINAIS and new_status not in STATUS_TERMINAIS:
                    old_closed = current.get("encerrado_em")
                    ws.cell(row, SHEETS["Chamados"].index("encerrado_em") + 1).value = None
                    self._append_history(wb, ticket_id, actor, "ATUALIZACAO", "encerrado_em", old_closed, None, "Reabertura do chamado.")
                ws.cell(row, SHEETS["Chamados"].index("atualizado_em") + 1).value = now
                ws.cell(row, SHEETS["Chamados"].index("atualizado_por") + 1).value = actor
                self._append_history(wb, ticket_id, actor, "ATUALIZACAO", "atualizado_em", current.get("atualizado_em"), now, comment)
                self._atomic_save(wb)
                wb.close()
                return len(changes)
        except Timeout as exc:
            raise RuntimeError("A base está ocupada por outra gravação. Tente novamente em alguns segundos.") from exc

    def record_handover(self, ticket_ids: list[str], source_shift: str, target_shift: str,
                        operator: str, summary: str, received_by: str = "") -> str:
        actor = str(operator).strip()
        ids = list(dict.fromkeys(str(x).strip() for x in ticket_ids if str(x).strip()))
        if not actor:
            raise ValueError("Informe o operador responsável pela passagem.")
        if not ids:
            raise ValueError("Selecione ao menos um chamado para transferir.")
        if source_shift == target_shift:
            raise ValueError("O turno de origem e o turno de destino devem ser diferentes.")
        if not summary.strip():
            raise ValueError("Descreva o resumo da passagem de turno.")
        try:
            with FileLock(str(self.lock_path), timeout=10):
                wb = self._load()
                ws = wb["Chamados"]
                selected: list[tuple[str, int, dict[str, Any]]] = []
                for ticket_id in ids:
                    row = self._find_ticket_row(ws, ticket_id)
                    if not row:
                        wb.close()
                        raise ValueError(f"Chamado {ticket_id} não foi encontrado.")
                    rec = self._row_map(ws, row)
                    if rec.get("status") in STATUS_TERMINAIS:
                        wb.close()
                        raise ValueError(f"O chamado {ticket_id} está encerrado e não pode ser transferido como pendência.")
                    selected.append((ticket_id, row, rec))
                self._backup_once_per_day()
                now = now_local()
                passage_id = f"PAS-{now:%Y%m%d%H%M%S}-{uuid.uuid4().hex[:5].upper()}"
                received = received_by.strip() or ""
                received_at = now if received else None
                ws_pass = wb["Passagens"]
                ws_pass.append([
                    passage_id, now, _safe_excel_value(source_shift), _safe_excel_value(target_shift),
                    _safe_excel_value(actor), len(selected), _safe_excel_value(", ".join(ids)),
                    _safe_excel_value(summary.strip()), _safe_excel_value(received), received_at,
                ])
                for ticket_id, row, rec in selected:
                    shift_col = SHEETS["Chamados"].index("turno_atual") + 1
                    old_shift = rec.get("turno_atual")
                    ws.cell(row, shift_col).value = _safe_excel_value(target_shift)
                    ws.cell(row, SHEETS["Chamados"].index("atualizado_em") + 1).value = now
                    ws.cell(row, SHEETS["Chamados"].index("atualizado_por") + 1).value = actor
                    self._append_history(wb, ticket_id, actor, "PASSAGEM", "turno_atual", old_shift, target_shift,
                                         f"{passage_id}: {summary.strip()}")
                self._atomic_save(wb)
                wb.close()
                return passage_id
        except Timeout as exc:
            raise RuntimeError("A base está ocupada por outra gravação. Tente novamente em alguns segundos.") from exc
